from pathlib import Path
import json,math,collections,hashlib,datetime
import numpy as np
import openpyxl
B=Path('/mnt/BioDisk/wangmingyuan/GI_preanalysis_20260923/CGNL1_phosphoproteome_qualification_20261018_v1');A=B/'audit/CGNL1_FIRST_REVIEW_20261022_v1';S=Path('/mnt/BioData/bio/wangmingyuan/CGNL1_phosphosite_bulk_576/source_search_593')
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def save(n,x):(A/n).write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')
def sid(v):return str(int(v)) if isinstance(v,(int,float)) else str(v).strip()
def numeric(v):
 try:return math.isfinite(float(v))
 except (ValueError,TypeError):return False
w=openpyxl.load_workbook(S/'B02_TableS1.xlsx',read_only=True,data_only=True);rows=list(w['S1A. Clinical info'].values);headers=list(rows[1]);clinical={sid(r[0]):dict(zip(headers,r)) for r in rows[2:] if r[0] is not None}
ex=B/'results/source_search_593/B02';prot=next(s for s in read(ex/'B02_protein_NODE.xlsx.extraction.json') if s['sheet']=='proteomics_NA');phos=next(s for s in read(ex/'B02_phospho_NODE.xlsx.extraction.json') if s['sheet']=='phospho_NA')
ids=[sid(v) for v in prot['headers'][0][2:]];pvals=dict(zip(ids,prot['hits'][0]['values'][2:]));svals=dict(zip([sid(v) for v in phos['headers'][0][12:]],next(r['values'][12:] for r in phos['hits'] if r['values'][0]=='CGNL1:S283')))
gw=openpyxl.load_workbook(S/'B02_groups_NODE.xlsx',read_only=True,data_only=True);batch=collections.defaultdict(set);channels=[]
for sh in gw:
 for row_number,r in enumerate(sh.iter_rows(values_only=True),1):
  if len(r)>2 and isinstance(r[1],str) and r[1].startswith('iCCA_'):
   for ci,v in enumerate(r[2:12],2):
    if isinstance(v,(int,float)):batch[sid(v)].add(r[1]);channels.append(dict(case=sid(v),batch=r[1],source_row=row_number,source_column=ci+1))
mapping_issue=dict(multiple_batch_ids={k:sorted(batch[k]) for k in ids if len(batch[k])>1},missing_batch_ids=[k for k in ids if not batch[k]],design_ids_absent_from_final_matrix=sorted(set(batch)-set(ids)),total_tumor_channels=len(channels),unique_design_case_ids=len(batch),final_unique_matrix_cases=len(set(ids)),interpretation='Repeated case IDs in TMT design are not independent patients; source final matrix has one column per case. No arbitrary batch assignment made')
keys=['Vascular invasion','Age','Sex','Tumor_size_diameter (cm)','Intrahepatic metastasis','Regional lymph Node metastasis','Distal metastasis'];domains={k:dict(collections.Counter(str(clinical[x].get(k)) for x in ids)) for k in keys}
records=[];excluded=[]
for ident in ids:
 r=clinical[ident];y=r.get('Vascular invasion');age=r.get('Age');sex=r.get('Sex');size=r.get('Tumor_size_diameter (cm)')
 good=y in ['Yes','No'] and numeric(age) and sex in ['Male','Female'] and numeric(size) and float(size)>0 and numeric(pvals[ident]) and numeric(svals[ident])
 if good:records.append(dict(case=ident,y=int(y=='Yes'),age=float(age),sex=int(sex=='Female'),log2_size=math.log2(float(size)),protein=float(pvals[ident]),site=float(svals[ident]),batch=next(iter(batch[ident])) if len(batch[ident])==1 else None))
 else:excluded.append(dict(case=ident,values={k:r.get(k) for k in keys[:4]}))
X=np.array([[1,x['site'],x['protein'],x['age'],x['sex'],x['log2_size']] for x in records]);br=[r for r in records if r['batch'] is not None];batches=sorted({r['batch'] for r in br});bx=np.array([[1,x['site'],x['protein'],x['age'],x['sex'],x['log2_size']] for x in br]);Xd=np.column_stack([bx]+[np.array([int(r['batch']==bb) for r in br]) for bb in batches[1:]])
ct={bb:{str(y):sum(r['batch']==bb and r['y']==y for r in br) for y in [0,1]} for bb in batches};yes=sum(r['y'] for r in records);no=len(records)-yes;byes=sum(r['y'] for r in br);bno=len(br)-byes
b02=dict(matrix_patients=len(ids),unique_clinical_registry=len(clinical),clinical_header=headers,domains=domains,complete_model_input=len(records),excluded=excluded,event_yes=yes,event_no=no,primary_nonintercept_parameters=5,minor_class_per_primary_parameter=min(yes,no)/5,design_rank=int(np.linalg.matrix_rank(X)),design_columns=X.shape[1],batch_mapping=mapping_issue,unambiguous_batch_diagnostic_subset=len(br),batch_count=len(batches),batch_fixed_effect_rank=int(np.linalg.matrix_rank(Xd)),batch_fixed_effect_columns=Xd.shape[1],minor_class_per_batch_fixed_effect_parameter=min(byes,bno)/(Xd.shape[1]-1),batch_outcome_counts=ct,single_outcome_batches=[bb for bb,v in ct.items() if min(v.values())==0],purity_like_fields=[h for h in headers if h and any(t in str(h).lower() for t in ['purity','stromal','immune'])],source_site_protein_pearson=float(np.corrcoef([r['site'] for r in records],[r['protein'] for r in records])[0,1]),clinical_effect_not_tested=True,diagnostic_subset_not_a_new_analysis_cohort=True)
save('B02_FIELD_AND_DESIGN_AUDIT.json',b02);save('B02_TMT_CHANNEL_MAP.json',channels)
# Check original baseline hashes, including all frozen plans and prior results.
before=read(A/'BASELINE_HASHES_BEFORE.json');after={}
for name in before:
 h=hashlib.sha256()
 with Path(name).open('rb') as f:
  for chunk in iter(lambda:f.read(8*1024*1024),b''):h.update(chunk)
 after[name]=h.hexdigest()
assert before==after
save('READ_ONLY_ACCEPTANCE.json',dict(utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),baseline_files=len(before),baseline_unchanged=True,new_clinical_effect_models=0,new_inferential_tests=0,initial_audit_stage='B02 unique-batch assumption rejected by source; continued with explicit ambiguity inventory',output_files=[p.name for p in A.glob('*.json')]))
print(json.dumps({k:v for k,v in b02.items() if k not in ['domains','clinical_header','batch_outcome_counts','excluded']}));print('READ_ONLY_ACCEPTANCE PASS')
