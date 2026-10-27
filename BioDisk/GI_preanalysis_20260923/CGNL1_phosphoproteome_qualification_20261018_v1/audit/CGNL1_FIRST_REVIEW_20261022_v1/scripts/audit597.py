from pathlib import Path
import json,csv,re,math,hashlib,datetime,collections,statistics
import numpy as np
import openpyxl
B=Path('/mnt/BioDisk/wangmingyuan/GI_preanalysis_20260923/CGNL1_phosphoproteome_qualification_20261018_v1')
D=Path('/mnt/BioData/bio/wangmingyuan/CGNL1_phosphosite_bulk_576')
A=B/'audit/CGNL1_FIRST_REVIEW_20261022_v1';A.mkdir(parents=True,exist_ok=True)
csv.field_size_limit(10000000)
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def save(n,x):
 p=A/n;p.write_text(json.dumps(x,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8');return p
def numeric(x):
 try:y=float(x)
 except (ValueError,TypeError):return False
 return math.isfinite(y)
def summary(v):
 x=np.asarray([z for z in v if numeric(z)],dtype=float)
 return dict(n=len(x),mean=float(x.mean()),median=float(np.median(x)),min=float(x.min()),max=float(x.max())) if len(x) else dict(n=0)
def sha(p):
 h=hashlib.sha256()
 with p.open('rb') as f:
  for c in iter(lambda:f.read(8*1024*1024),b''):h.update(c)
 return h.hexdigest()
protected=[B.parent/'CGNL1_CURRENT_STATE_20261017.json',B.parent/'CGNL1_MASTER_PLAN_CURRENT.json']
for folder in ['results/H01_582','results/H02_recovered_590','results/B03_recovered_590','results/PDAC_phosphosite_357','results/HCC_exploratory_572','planning/CGNL1_MASTER_PLAN_20261022_v1']:
 protected.extend(p for p in (B/folder).glob('*') if p.is_file())
before={str(p):sha(p) for p in protected};save('BASELINE_HASHES_BEFORE.json',before)

# Directly audit the matrices and saved offsets; do not run original writers.
def matrix_audit(path,cohort):
 with path.open(encoding='utf-8') as f:
  reader=csv.DictReader(f,delimiter='\t');heads=reader.fieldnames
  cols=[c for c in heads if re.fullmatch(r'[TP]\d+ Log Ratio',c)] if cohort=='H01' else [c for c in heads if c.startswith('Log ')]
  mat=[];target=[];special={}
  for row in reader:
   key=row[heads[0]]
   vals=[float(row[c]) if numeric(row[c]) else np.nan for c in cols]
   if key.lower() in ['mean','median','stddev']:
    special[key]=dict(zip(cols,[float(x) if np.isfinite(x) else None for x in vals]));continue
   mat.append(vals)
   if row.get('Gene')=='CGNL1':target.append(row)
 mat=np.asarray(mat);median=np.nanmedian(mat,axis=0)
 out=dict(source=str(path),rows=len(mat),columns=len(cols),source_special_rows=list(special),computed_column_medians=summary(median),column_medians=dict(zip(cols,median.tolist())),special_rows=special,target_row_count=len(target))
 return out,target,cols
hm={};ht={}
for kind,name in [('protein','Zhou_Liver_Cancer_Proteome.tmt11.tsv'),('phosphosite','Zhou_Liver_Cancer_Phosphoproteome.phosphosite.tmt11.tsv')]:
 hm[kind],ht[kind],_=matrix_audit(D/'H01'/name,'H01')
offset=read(B/'results/H01_582/NORMALIZATION.json')
for kind in hm:hm[kind]['max_saved_offset_error']=max(abs(v-offset[kind][k]) for k,v in hm[kind]['column_medians'].items())
pairs=read(B/'results/H01_582/ANALYSIS_PAIRING.json');values=[x for x in read(B/'results/H01_582/PATIENT_VALUES.json') if x['site']=='S283']
decomp=[]
for row in values:
 t,n=row['tumor']+' Log Ratio',row['normal']+' Log Ratio';ps=hm['protein']['column_medians'];ss=hm['phosphosite']['column_medians']
 correction=(ss[t]-ss[n])-(ps[t]-ps[n]);raw=row['unscaled_adjusted'];center=row['adjusted'];assert abs(raw-correction-center)<1e-12
 decomp.append(dict(case=row['case'],source_adjusted=raw,centered_adjusted=center,site_median_change=ss[t]-ss[n],protein_median_change=ps[t]-ps[n],net_subtracted_offset=correction,sign_flipped=(raw*center<0)))
h01=dict(matrices=hm,S283_decomposition=dict(n=len(decomp),source_adjusted=summary([r['source_adjusted'] for r in decomp]),centered_adjusted=summary([r['centered_adjusted'] for r in decomp]),site_median_change=summary([r['site_median_change'] for r in decomp]),protein_median_change=summary([r['protein_median_change'] for r in decomp]),net_subtracted_offset=summary([r['net_subtracted_offset'] for r in decomp]),sign_flips=sum(r['sign_flipped'] for r in decomp),source_positive=sum(r['source_adjusted']>0 for r in decomp),centered_positive=sum(r['centered_adjusted']>0 for r in decomp)),pairing_coverage=read(B/'results/H01_582/PAIRING_COVERAGE.json'),interpretation='Arithmetic reconstruction only; no primary normalization substituted')
save('H01_NORMALIZATION_AUDIT.json',h01);save('H01_PATIENT_COMPONENTS.json',decomp);print('H01',json.dumps(h01['S283_decomposition']),flush=True)
bm={}
for kind,name in [('protein','iCC_NCC_Proteome.tmt10.tsv'),('phosphosite','iCC_NCC_Phosphoproteome.phosphosite.tmt10.tsv')]:bm[kind],_,_=matrix_audit(D/'B03'/name,'B03')
save('B03_NORMALIZATION_AUDIT.json',bm)

# Source-level PDAC component verification, without new P-value testing.
raws={}
for kind in ['proteomics','phospho']:
 for condition in ['Tumor','Normal']:
  p=next((B/'raw').glob('PDAC_'+kind+'*_'+condition+'.txt'))
  with p.open() as f:
   r=csv.reader(f,delimiter='\t');h=next(r);rows={}
   for v in r:
    if v[0].split('|')[0].split('.')[0]=='ENSG00000128849':rows[v[0]]=dict(zip(h[1:],[float(x) if numeric(x) else None for x in v[1:]]))
  raws[kind,condition]=rows
with (B/'results/PDAC_phosphosite_357/patient_site_values.tsv').open() as f:pdacrows=[r for r in csv.DictReader(f,delimiter='\t') if '|S283|' in r['site']]
pdac=[];maxerr=0
for r in pdacrows:
 case=r['case'];site=r['site'];prot=next(iter(raws['proteomics','Tumor']))
 st=raws['phospho','Tumor'][site][case];sn=raws['phospho','Normal'][site][case];pt=raws['proteomics','Tumor'][prot][case];pn=raws['proteomics','Normal'][prot][case]
 sd=st-sn;pd=pt-pn;ad=sd-pd;maxerr=max(maxerr,abs(ad-float(r['adjusted_delta'])))
 pdac.append(dict(case=case,site_delta=sd,protein_delta=pd,adjusted_delta=ad))
pdacsummary=dict(n=len(pdac),max_source_reconstruction_error=maxerr,site_change=summary([r['site_delta'] for r in pdac]),protein_change=summary([r['protein_delta'] for r in pdac]),adjusted_change=summary([r['adjusted_delta'] for r in pdac]),adjusted_positive=sum(r['adjusted_delta']>0 for r in pdac),adjusted_positive_but_site_negative=sum(r['adjusted_delta']>0 and r['site_delta']<0 for r in pdac),adjusted_positive_both_components_negative=sum(r['adjusted_delta']>0 and r['site_delta']<0 and r['protein_delta']<0 for r in pdac),not_site_occupancy=True)
save('PDAC_COMPONENT_AUDIT.json',pdacsummary);save('PDAC_PATIENT_COMPONENTS.json',pdac);print('PDAC',json.dumps(pdacsummary),flush=True)

# Reconcile H02 author maps and the missing complete cases.
H=D/'H02';w=openpyxl.load_workbook(next((H/'global_589').rglob('*.xlsx')),read_only=True,data_only=True)
maps={};stage_maps={}
for sheet in w:
 if sheet.title=='Description':continue
 mp={};st={}
 for row in sheet.iter_rows(min_row=2,values_only=True):
  if len(row)<5 or not row[2] or not row[4]:continue
  key=str(row[2]).replace('_','').strip();mp[key]=str(row[4]).strip();st[key]=str(row[3]).strip()
 maps[sheet.title]=mp;stage_maps[sheet.title]=st
qname=next(k for k in maps if k.startswith('Quality'));nname=next(k for k in maps if k.startswith('Quantitative'))
qmap,nmap=maps[qname],maps[nname];nst=stage_maps[nname]
ph=read(B/'results/H02_source_587/Phospho (STY)Sites.txt.target.json');s283=next(r for r in ph['rows'] if r['Amino acid']=='S' and r['Position']=='283')
rawcases=sorted({h[10:-2] for h in ph['headers'] if re.fullmatch('Intensity [0-9]+_T',h) and h[:-1]+'P' in ph['headers']})
gp=next((H/'global_589').rglob('proteinGroups.txt'))
with gp.open() as f:
 rr=csv.DictReader(f,delimiter='\t');gcols=rr.fieldnames;grows=[r for r in rr if 'CGNL1' in r.get('Gene names','').split(';')]
assert len(grows)==1;grow=grows[0];reconciliation=[];coverage=[]
for raw in rawcases:
 t,n=raw+'T',raw+'P';rec=dict(raw=raw)
 if t not in qmap or n not in qmap:reason='no_authoritative_raw_id_crosswalk'
 elif t not in nmap or n not in nmap:reason='not_in_author_quantitative_set'
 elif nst[t] not in ['A','0'] or nst[n] not in ['A','0']:reason='not_BCLC_A_or_0'
 elif 'Intensity '+t not in gcols or 'Intensity '+n not in gcols:reason='missing_independent_total_protein_column'
 else:reason='eligible'
 rec.update(reason=reason,quality_case_T=qmap.get(t),quantitative_case_T=nmap.get(t),stage=nst.get(t));reconciliation.append(rec)
 if reason!='eligible':continue
 assert nmap[t].endswith('_T') and nmap[n].endswith('_N') and nmap[t][:-2]==nmap[n][:-2]
 def positive(v):return numeric(v) and float(v)>0
 fields={'site_T':s283.get('Intensity '+raw+'_T'),'site_N':s283.get('Intensity '+raw+'_P'),'protein_T':grow.get('Intensity '+t),'protein_N':grow.get('Intensity '+n)}
 missing=[k for k,v in fields.items() if not positive(v)]
 probs={k:s283.get('Localization prob '+raw+'_'+suffix) for k,suffix in [('localization_T','T'),('localization_N','P')]}
 coverage.append(dict(raw=raw,case=nmap[t][:-2],complete=not missing,missing=missing,values={k:float(v) if numeric(v) else None for k,v in fields.items()},**{k:float(v) if numeric(v) else None for k,v in probs.items()}))
eligible=read(B/'results/H02_recovered_590/PAIRING.json')['eligible'];assert {r['raw'] for r in coverage}=={r['raw'] for r in eligible}
frozen=[r for r in read(B/'results/H02_recovered_590/PATIENT_VALUES.json') if r['site']=='S283'];assert {r['case'] for r in coverage if r['complete']}=={r['case'] for r in frozen}
effects=[r['adjusted'] for r in frozen];unc=[r['uncentered'] for r in frozen];total=sum(effects);lopo=[(total-x)/(len(effects)-1) for x in effects]
groups=collections.Counter(r['reason'] for r in reconciliation);missing=collections.Counter(k for r in coverage for k in r['missing'])
lowloc=[dict(case=r['case'],which=k,value=r[k]) for r in coverage if r['complete'] for k in ['localization_T','localization_N'] if r[k] is None or r[k]<.75]
clinical=H/'41586_2019_987_MOESM5_ESM.xlsx';clinical_summary={}
cw=openpyxl.load_workbook(clinical,read_only=True,data_only=True)
for sheet in cw:
 rows=list(sheet.values)
 if rows and rows[0] and rows[0][0]=='ID':
  mp={str(r[0]):r for r in rows[1:] if r[0]};head=rows[0]
  for field in ['MVI','AFP']:
   if field not in head:continue
   ix=head.index(field);counts={}
   for complete in [True,False]:counts[str(complete)]=dict(collections.Counter(str(mp[r['case']][ix]) if r['case'] in mp and mp[r['case']][ix] is not None else 'missing' for r in coverage if r['complete']==complete))
   clinical_summary[field]=counts
h02=dict(raw_pair_ids=len(rawcases),reason_counts=dict(groups),eligible=len(coverage),complete_S283=len(frozen),missing_field_counts=dict(missing),global_S283_localization_probability=float(s283['Localization prob']),complete_pairs_with_low_or_absent_per_sample_localization=lowloc,leave_one_out_means=summary(lopo),primary_negative=sum(x<0 for x in effects),uncentered_negative=sum(x<0 for x in unc),individual_sign_flips=sum(x*y<0 for x,y in zip(effects,unc)),clinical_missingness_counts=clinical_summary,clinical_comparison='Descriptive source-coded MVI/AFP counts only; no new inferential missingness test; other clinical PDF fields not examined',MBR_metadata={p.name+'_'+str(i):'Match between runs\tTrue' in p.read_text() for i,p in enumerate([H/'selected_585/parameters.txt',next((H/'global_589').rglob('parameters.txt'))])})
save('H02_EXCLUSION_LEDGER.json',reconciliation);save('H02_S283_COVERAGE.json',coverage);save('H02_ROBUSTNESS_AUDIT.json',h02);print('H02',json.dumps(h02),flush=True)

# B02 inputs and design feasibility only. No clinical target model is fitted.
S=D/'source_search_593';w=openpyxl.load_workbook(S/'B02_TableS1.xlsx',read_only=True,data_only=True);rows=list(w['S1A. Clinical info'].values);headers=list(rows[1]);clinical={str(r[0]).strip():dict(zip(headers,r)) for r in rows[2:] if r[0] is not None}
ex=B/'results/source_search_593/B02'
prot=next(s for s in read(ex/'B02_protein_NODE.xlsx.extraction.json') if s['sheet']=='proteomics_NA');phos=next(s for s in read(ex/'B02_phospho_NODE.xlsx.extraction.json') if s['sheet']=='phospho_NA')
ids=[str(int(v)) for v in prot['headers'][0][2:]];pvals=dict(zip(ids,prot['hits'][0]['values'][2:]));svals=dict(zip([str(int(v)) for v in phos['headers'][0][12:]],next(r['values'][12:] for r in phos['hits'] if r['values'][0]=='CGNL1:S283')))
gw=openpyxl.load_workbook(S/'B02_groups_NODE.xlsx',read_only=True,data_only=True);batch=collections.defaultdict(set)
for sh in gw:
 for r in sh.iter_rows(values_only=True):
  if len(r)>2 and isinstance(r[1],str) and r[1].startswith('iCCA_'):
   for v in r[2:12]:
    if isinstance(v,(int,float)):batch[str(int(v))].add(r[1])
assert all(len(batch[k])==1 for k in ids)
keys=['Vascular invasion','Age','Sex','Tumor_size_diameter (cm)','Intrahepatic metastasis','Regional lymph Node metastasis','Distal metastasis']
domains={k:dict(collections.Counter(str(clinical[x].get(k)) for x in ids)) for k in keys}
records=[];excluded=[]
for ident in ids:
 r=clinical[ident];y=r.get('Vascular invasion');age=r.get('Age');sex=r.get('Sex');size=r.get('Tumor_size_diameter (cm)')
 good=y in ['Yes','No'] and numeric(age) and sex in ['Male','Female'] and numeric(size) and float(size)>0 and numeric(pvals[ident]) and numeric(svals[ident])
 if good:records.append(dict(case=ident,y=int(y=='Yes'),age=float(age),sex=int(sex=='Female'),log2_size=math.log2(float(size)),protein=float(pvals[ident]),site=float(svals[ident]),batch=next(iter(batch[ident]))))
 else:excluded.append(dict(case=ident,values={k:r.get(k) for k in keys[:4]}))
X=np.array([[1,x['site'],x['protein'],x['age'],x['sex'],x['log2_size']] for x in records]);batches=sorted({r['batch'] for r in records});Xd=np.column_stack([X]+[np.array([int(r['batch']==bb) for r in records]) for bb in batches[1:]])
ct={bb:{str(y):sum(r['batch']==bb and r['y']==y for r in records) for y in [0,1]} for bb in batches};yes=sum(r['y'] for r in records);no=len(records)-yes
b02=dict(matrix_patients=len(ids),unique_clinical_registry=len(clinical),clinical_header=headers,domains=domains,complete_model_input=len(records),excluded=excluded,event_yes=yes,event_no=no,primary_nonintercept_parameters=5,minor_class_per_primary_parameter=min(yes,no)/5,design_rank=int(np.linalg.matrix_rank(X)),design_columns=X.shape[1],batch_count=len(batches),batch_fixed_effect_rank=int(np.linalg.matrix_rank(Xd)),batch_fixed_effect_columns=Xd.shape[1],minor_class_per_batch_fixed_effect_parameter=min(yes,no)/(Xd.shape[1]-1),batch_outcome_counts=ct,single_outcome_batches=[bb for bb,v in ct.items() if min(v.values())==0],missing_purity_fields=[h for h in headers if h and any(t in str(h).lower() for t in ['purity','stromal','immune'])],source_site_protein_pearson=float(np.corrcoef([r['site'] for r in records],[r['protein'] for r in records])[0,1]),clinical_effect_not_tested=True)
save('B02_FIELD_AND_DESIGN_AUDIT.json',b02);print('B02',json.dumps({k:v for k,v in b02.items() if k not in ['domains','clinical_header','batch_outcome_counts','excluded']}),flush=True)
after={str(p):sha(p) for p in protected};assert before==after
save('READ_ONLY_ACCEPTANCE.json',dict(utc=datetime.datetime.now(datetime.timezone.utc).isoformat(),baseline_files=len(protected),baseline_unchanged=True,new_clinical_effect_models=0,new_inferential_tests=0,output_files=[p.name for p in A.glob('*.json')]))
print('READ_ONLY_ACCEPTANCE PASS',flush=True)
