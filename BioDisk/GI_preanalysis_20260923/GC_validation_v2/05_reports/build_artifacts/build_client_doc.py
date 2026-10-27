import re
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(r'C:\Users\31937\GC_validation_20260923');D=R/'delivery';p=D/'客户版研究建议书.md';text=p.read_text(encoding='utf-8-sig');text=text.replace('1. Charton等 胃癌原发与腹膜转移多组学研究','1. 胃癌原发与腹膜转移多组学研究');text=text.replace('说明结果不是由单个患者造成。','说明结果不是由单个患者造成。[1]');text=text.replace('这是同一细胞类型内的探索性线索。','这是同一细胞类型内的探索性线索。[8]');text=text.replace('3个预设定位终点内FDR均为1.12×10⁻⁸。','3个预设定位终点内FDR均为1.12×10⁻⁸。[3]');text=text.split('\n8. GSE234129')[0]+'\n8. GSE234129 Evolution of immune and stromal cell states and ecotypes during gastric adenocarcinoma progression GEO公开处理矩阵。\n';text=re.sub(r'(\[\d+\])(?:\s*\1)+',r'\1',text);p.write_text(text,encoding='utf-8')
doc=Document();s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.75);s.left_margin=s.right_margin=Inches(.8);s.footer_distance=Inches(.3)
styles=doc.styles
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
 st=styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0);rpr=st.element.get_or_add_rPr();fonts=rpr.find(qn('w:rFonts'))
 if fonts is None:fonts=OxmlElement('w:rFonts');rpr.append(fonts)
 fonts.set(qn('w:eastAsia'),'宋体' if name in ['Normal','Caption'] else '微软雅黑');fonts.set(qn('w:ascii'),'Calibri');fonts.set(qn('w:hAnsi'),'Calibri')
styles['Normal'].font.size=Pt(11);styles['Normal'].paragraph_format.line_spacing=1.18;styles['Normal'].paragraph_format.space_after=Pt(8)
styles['Title'].font.size=Pt(22);styles['Title'].font.bold=True;styles['Title'].paragraph_format.space_after=Pt(10)
styles['Heading 1'].font.size=Pt(15);styles['Heading 1'].paragraph_format.space_before=Pt(5);styles['Heading 1'].paragraph_format.space_after=Pt(10);styles['Heading 1'].paragraph_format.keep_with_next=True
styles['Heading 2'].font.size=Pt(12);styles['Heading 2'].paragraph_format.space_after=Pt(7)
styles['Caption'].font.size=Pt(9.5);styles['Caption'].font.bold=False;styles['Caption'].paragraph_format.line_spacing=1.12;styles['Caption'].paragraph_format.space_after=Pt(9)
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=footer.add_run('第 ');r.font.size=Pt(9);fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld);footer.add_run(' 页').font.size=Pt(9)
# Set all table borders explicitly, keep rows intact and headers repeated.
def table(rows):
 t=doc.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=[2.95,3.95] if len(rows[0])==2 else [1.45,2.55,2.9]
 for i,w in enumerate(widths):t.columns[i].width=Inches(w)
 pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
 for side in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
 pr.append(b)
 for ri,row in enumerate(rows):
  cells=t.rows[0].cells if ri==0 else t.add_row().cells
  trpr=cells[0]._tc.getparent().get_or_add_trPr();ns=OxmlElement('w:cantSplit');trpr.append(ns)
  if ri==0:rep=OxmlElement('w:tblHeader');trpr.append(rep)
  for ci,txt in enumerate(row):
   c=cells[ci];c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcpr=c._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
   for k,val in [('top','95'),('bottom','95'),('left','100'),('right','100')]:e=OxmlElement('w:'+k);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margin.append(e)
   tcpr.append(margin);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'244D66' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));tcpr.append(sh)
   pp=c.paragraphs[0];pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.12;rr=pp.add_run(txt);rr.font.size=Pt(10.5);rr.font.bold=ri==0
   if ri==0:rr.font.color.rgb=RGBColor(255,255,255)
   if ri==0 or (ci==0 and len(txt)<=12):pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
refs={1:'https://doi.org/10.1038/s42003-026-09906-2',2:'https://doi.org/10.1038/s41467-023-36310-9',3:'https://doi.org/10.3390/ijms25074117',4:'https://doi.org/10.1136/gutjnl-2024-333617',5:'https://doi.org/10.1136/jitc-2026-015464',6:'https://pubmed.ncbi.nlm.nih.gov/35198079/',7:'https://pubmed.ncbi.nlm.nih.gov/42597715/',8:'https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSE234129'}
def hyperlink(p,text,url):
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True));r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'175C7F');rp.append(color);r.append(rp);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
def figure(name,caption,width=6.65):
 pp=doc.add_paragraph();pp.paragraph_format.space_before=Pt(4);pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.keep_with_next=True;pp.add_run().add_picture(str(R/'figures'/name),width=Inches(width));cp=doc.add_paragraph(caption,style='Caption');cp.paragraph_format.keep_with_next=False
lines=text.splitlines();section=0;i=0;inrefs=False
while i<len(lines):
 line=lines[i].strip()
 if not line:i+=1;continue
 if line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
 elif line.startswith('## '):
  if section==2:figure('F1_tissue_source.png','图1 组织层差异及敏感性分析。左图为同一17对患者的模型比较，右图为配对变化相关性；均不能替代细胞内机制验证。',6.8)
  section+=1
  if section in [2,3,4,7]:doc.add_page_break()
  doc.add_paragraph(line[3:],style='Heading 1');inrefs=False
 elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2');inrefs=True
 elif line.startswith('|'):
  if section==5:
   doc.add_page_break();doc.add_paragraph('数据分层与验证资源',style='Heading 1')
  rows=[]
  while i<len(lines) and lines[i].strip().startswith('|'):
   row=[a.strip() for a in lines[i].strip().strip('|').split('|')]
   if not all(set(a)<=set('-: ') for a in row):rows.append(row)
   i+=1
  table(rows);continue
 else:
  pp=doc.add_paragraph(line)
  if line=='2026年9月23日':pp.paragraph_format.space_after=Pt(12);pp.runs[0].font.size=Pt(10);pp.runs[0].font.color.rgb=RGBColor.from_string('555555')
  if inrefs:
   pp.paragraph_format.space_after=Pt(6);pp.paragraph_format.line_spacing=1.08
   for rr in pp.runs:rr.font.size=Pt(9.5)
   num=int(line.split('.')[0]) if line.split('.')[0].isdigit() else None
   if num in refs:hyperlink(pp,'  原文',refs[num])
  if section==3 and line.startswith('在4位'):figure('F2_paired_macrophages.png','图2 同患者同细胞类型的探索性比较。每条线代表1位患者，n=4；未以细胞数扩大统计样本量。',6.25)
  if section==3 and line.startswith('独立GSE228598'):figure('F3_independent_localization.png','图3 独立28例的患者内细胞定位。每条线代表1位患者；C1Q巨噬细胞富集规则未使用CD36、TREM2或FABP4。',6.25)
 i+=1
doc.core_properties.title='胃癌腹膜转移脂质相关髓系微环境研究建议书';doc.core_properties.subject='预验证证据与正式研究启动方案';doc.core_properties.author='';doc.core_properties.keywords='胃癌 腹膜转移 多组学 髓系微环境';out=D/'胃癌腹膜转移研究建议书.docx';
for root in [doc.element,doc.styles.element]:
 for node in root.xpath('.//w:pBdr'):
  node.getparent().remove(node)
doc.save(out);print(out)
